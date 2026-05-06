#!/usr/bin/env python3
"""
BOM 成本计算脚本
计算指定 SKU 的小包装/大包装/总成本（按 1kg 和单品两个维度）
"""
import argparse
import openpyxl
from pathlib import Path

BOM_DIR = Path("/Users/taojing/Documents/OpenClaw_Docs/Bom/")


def get_bom_path(month: str) -> Path:
    year, month_num = month.split("-")
    file_name = f"{year}年{int(month_num)}月bom表.xlsx"
    path = BOM_DIR / file_name
    if not path.exists():
        raise FileNotFoundError(f"BOM 文件不存在: {path}")
    return path


def load_bom(path: Path):
    return openpyxl.load_workbook(path, data_only=True)


def parse_product_info(sheet):
    """解析 Row 1-4 的产品维度信息（列从 Col 3 开始，每产品占2列）"""
    products = []
    max_col = sheet.max_column
    col = 3  # 产品从第3列开始
    
    while col + 1 <= max_col:
        code = sheet.cell(row=1, column=col).value
        name = sheet.cell(row=2, column=col).value
        spec = sheet.cell(row=3, column=col).value
        weight = sheet.cell(row=4, column=col).value
        
        if not code or (isinstance(code, str) and not code.strip()):
            break
        
        products.append({
            "code": str(code).strip(),
            "name": str(name).strip() if name else "",
            "spec": str(spec).strip() if spec else "",
            "inbound_weight": float(weight) if weight and str(weight).strip() else 0,
            "qty_col": col,
            "amt_col": col + 1
        })
        col += 2
    
    return products


def find_sku(products, sku_code):
    """根据产品代码查找 SKU"""
    for p in products:
        if p["code"] == sku_code:
            return p
    return None


def calc_process_cost(sheet, sku_code, is_large=False):
    """计算某道工序的成本"""
    # 大包装 SKU 无后缀，小包装 SKU 有 "-1" 后缀
    if is_large:
        search_code = sku_code
    else:
        search_code = sku_code if sku_code.endswith("-1") else sku_code + "-1"
    
    products = parse_product_info(sheet)
    sku = find_sku(products, search_code)
    
    if not sku:
        raise ValueError(f"SKU {search_code} 在当前 Sheet 中不存在")
    
    # 工序总成本 = 该 SKU 所有物料金额之和
    total_cost = 0
    for row in range(6, sheet.max_row + 1):
        amt = sheet.cell(row=row, column=sku["amt_col"]).value
        if amt:
            total_cost += float(amt)
    
    return {
        "sku_code": search_code,
        "sku_name": sku["name"],
        "sku_spec": sku["spec"],
        "inbound_weight": sku["inbound_weight"],
        "process_cost": total_cost,
        "cost_per_kg": total_cost / sku["inbound_weight"] if sku["inbound_weight"] > 0 else 0
    }


def output_table(cost_data, process="both"):
    """输出飞书表格格式"""
    print("| 项目 | 数值 |")
    print("| --- | --- |")
    print(f"| 产品代码 | {cost_data['small']['sku_code'] if cost_data.get('small') else cost_data['large']['sku_code']} |")
    print(f"| 品名 | {cost_data['small']['sku_name'] if cost_data.get('small') else cost_data['large']['sku_name']} |")
    print(f"| 规格 | {cost_data['small']['sku_spec'] if cost_data.get('small') else cost_data['large']['sku_spec']} |")
    
    if process in ("small", "both"):
        print()
        print("**小包装工序**")
        print(f"| 入库重量 | {cost_data['small']['inbound_weight']:.2f} kg |")
        print(f"| 工序成本 | ¥{cost_data['small']['process_cost']:.2f} |")
        print(f"| 1kg 成本 | ¥{cost_data['small']['cost_per_kg']:.4f}/kg |")
    
    if process in ("large", "both"):
        print()
        print("**大包装工序**")
        print(f"| 入库重量 | {cost_data['large']['inbound_weight']:.2f} kg |")
        print(f"| 工序成本 | ¥{cost_data['large']['process_cost']:.2f} |")
        print(f"| 1kg 成本 | ¥{cost_data['large']['cost_per_kg']:.4f}/kg |")
    
    if process == "both" and cost_data.get("small") and cost_data.get("large"):
        total_per_kg = cost_data["small"]["cost_per_kg"] + cost_data["large"]["cost_per_kg"]
        spec_kg = parse_spec_to_kg(cost_data["small"]["sku_spec"])
        if spec_kg:
            unit_cost = total_per_kg * spec_kg
            print()
            print("**汇总**")
            print(f"| 总 1kg 成本 | ¥{total_per_kg:.4f}/kg |")
            print(f"| 规格(kg) | {spec_kg} kg |")
            print(f"| 单品成本 | ¥{unit_cost:.4f} |")
    
    print()
    print(f"数据来源：{cost_data['month']} BOM 表")


def parse_spec_to_kg(spec):
    """从规格字符串提取 kg 数值，如 '500g/盒' -> 0.5"""
    import re
    match = re.search(r'(\d+(?:\.\d+)?)\s*[gG]', spec)
    if match:
        return float(match.group(1)) / 1000
    match = re.search(r'(\d+(?:\.\d+)?)\s*[kK][gG]', spec)
    if match:
        return float(match.group(1))
    return None


def main():
    parser = argparse.ArgumentParser(description="计算 BOM 产品成本")
    parser.add_argument("--month", required=True, help="月份，如 2026-01")
    parser.add_argument("--sku", required=True, help="产品代码，如 YQ00109XS")
    parser.add_argument("--process", default="both", choices=["small", "large", "both"], help="工序")
    args = parser.parse_args()
    
    path = get_bom_path(args.month)
    wb = load_bom(path)
    
    result = {"month": args.month, "sku": args.sku}
    
    if args.process in ("small", "both"):
        result["small"] = calc_process_cost(wb.worksheets[0], args.sku, is_large=False)
    
    if args.process in ("large", "both"):
        result["large"] = calc_process_cost(wb.worksheets[1], args.sku, is_large=True)
    
    output_table(result, args.process)


if __name__ == "__main__":
    main()