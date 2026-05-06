#!/usr/bin/env python3
"""
BOM 表读取脚本
读取 xlsx 文件，解析小包装/大包装 sheet，返回结构化 JSON
"""
import argparse
import json
import openpyxl
from pathlib import Path

BOM_DIR = Path("/Users/taojing/Documents/OpenClaw_Docs/Bom/")


def get_bom_path(month: str) -> Path:
    """根据月份获取 BOM 文件路径"""
    # month 格式: 2026-01
    year, month_num = month.split("-")
    file_name = f"{year}年{int(month_num)}月bom表.xlsx"
    path = BOM_DIR / file_name
    if not path.exists():
        raise FileNotFoundError(f"BOM 文件不存在: {path}")
    return path


def load_bom(path: Path):
    """加载 BOM xlsx"""
    wb = openpyxl.load_workbook(path, data_only=True)
    return wb


def parse_product_info(sheet):
    """解析 Row 1-4 的产品维度信息
    
    实际列结构：
    - Col 1: 物料编码（标题列）
    - Col 2: 耗用材料（标题列）
    - Col 3+: 产品数据（每产品占2列：数量 + 金额）
    """
    products = []
    max_col = sheet.max_column
    
    # Row 5 是表头行：物料编码 | 耗用材料 | 数量(产品1) | 金额(产品1) | 数量(产品2) | 金额(产品2) | ...
    # 产品列从 Col 3 开始，每 2 列为一个产品
    col = 3
    while col + 1 <= max_col:
        code = sheet.cell(row=1, column=col).value
        name = sheet.cell(row=2, column=col).value
        spec = sheet.cell(row=3, column=col).value
        weight = sheet.cell(row=4, column=col).value
        
        if not code or (isinstance(code, str) and not code.strip()):
            break
        
        products.append({
            "code": str(code).strip(),
            "name": str(name).strip() if name else "",
            "spec": str(spec).strip() if spec else "",
            "inbound_weight": float(weight) if weight and str(weight).strip() else 0,
            "qty_col": col,      # 数量列
            "amt_col": col + 1   # 金额列
        })
        col += 2
    
    return products


def parse_material_consumption(sheet, products):
    """解析 Row 6+ 的物料消耗明细，返回以物料名称为 key 的字典"""
    materials = {}  # {material_name: {product_code: {"qty": x, "amt": y}}}
    
    for row in range(6, sheet.max_row + 1):
        material_code = sheet.cell(row=row, column=1).value
        material_name = sheet.cell(row=row, column=2).value
        
        if material_code is None or material_name is None:
            continue
        
        material_name = str(material_name).strip()
        
        for p in products:
            qty = sheet.cell(row=row, column=p["qty_col"]).value
            amt = sheet.cell(row=row, column=p["amt_col"]).value
            
            if qty is not None and str(qty).strip() != '':
                if material_name not in materials:
                    materials[material_name] = {}
                materials[material_name][p["code"]] = {
                    "qty": float(qty),
                    "amt": float(amt) if amt else 0
                }
    
    return materials


def sheet_to_json(sheet_name: str, month: str, sheet_idx: int = 0):
    """将指定 sheet 转为 JSON"""
    path = get_bom_path(month)
    wb = load_bom(path)
    
    if sheet_name == "small":
        sheet = wb.worksheets[0]
    elif sheet_name == "large":
        sheet = wb.worksheets[1]
    else:
        raise ValueError(f"Unknown sheet type: {sheet_name}")
    
    products = parse_product_info(sheet)
    materials = parse_material_consumption(sheet, products)
    
    return {
        "month": month,
        "sheet": sheet_name,
        "products": products,
        "materials": materials
    }


def output_table(data):
    """输出飞书表格格式"""
    print(f"| {'产品代码':<20} | {'品名':<15} | {'规格':<12} | {'入库重量(kg)':<12} |")
    print(f"| {'-'*20} | {'-'*15} | {'-'*12} | {'-'*12} |")
    
    for p in data["products"]:
        print(f"| {p['code']:<20} | {p['name']:<15} | {p['spec']:<12} | {p['inbound_weight']:<12} |")
    
    print()
    print(f"共 {len(data['products'])} 个产品，{len(data['materials'])} 种物料")
    print(f"数据来源：{data['month']} {'小包装' if data['sheet']=='small' else '大包装'} Sheet")


def main():
    parser = argparse.ArgumentParser(description="读取 BOM 表")
    parser.add_argument("--month", required=True, help="月份，如 2026-01")
    parser.add_argument("--sheet", required=True, choices=["small", "large"], help="小包装或大包装")
    args = parser.parse_args()
    
    data = sheet_to_json(args.sheet, args.month)
    output_table(data)


if __name__ == "__main__":
    main()